#Día 239 / 365
"""
Clasificador automático de correos (Spam vs No Spam)
"""

from sklearn.feature_extraction.text import CountVectorizer
from sklearn.naive_bayes import MultinomialNB

#Dataset pequeño de ejemplo
correo  = [

    "Gana dinero rápido desde casa",
    "Oferta exclusiva, compra ahora y recibe un descuento",
    "Hola amigo, ¿cómo estás?",
    "Recordatorio de la reunión de mañana",
    "Has sido seleccionado para un premio especial",
    "Nos vemos en el almuerzo de hoy",
    "Compra medicamentos sin receta aquí",
    "Tu pedido ha sido enviado con éxito"
]

etiquetas = [
    "spam", "spam", "no_spam", "no_spam", 
    "spam", "no_spam", "spam", "no_spam"

]

Vectorizador = CountVectorizer()
X = Vectorizador.fit_transform(correo)

#Entrenar modelo
modelo = MultinomialNB()
modelo.fit(X, etiquetas)

#Probar con un nuevo correo
nuvo_correo = input("Escrbie el texto del correo: ")
x_nuevo = Vectorizador.transform([nuvo_correo])
prediccion = modelo.predict(x_nuevo)[0]

print(f"El correo se clasifica como: {prediccion.upper()}")



# Día 240 / 365
"""
Dashboard de Productividad de Empleados
"""
import pandas as pd
import matplotlib.pyplot as plt

#Cargar datos de ejemplo
data = {
    "empleado": ["Ana", "Luis", "Carlos", "Marta", "Sofía"],
    "horas_trabajadas": [40, 38, 45, 36, 42],
    "tareas_completadas": [25, 20, 30, 18, 28],
    "proyectos_asignados": [30, 25, 35, 20, 32]
}

df = pd.DataFrame(data)

#Calcular métricas
df["eficiencia"] = df["tareas_completadas"] / df["proyectos_asignados"]

promedio_horas = df["horas_trabajadas"].mean()
print(f"Promedio de horas trabajadas: {promedio_horas:.2f}")

print("\nRanking por eficiencia:")
print(df.sort_values("eficiencia", ascending=False)[["empleado", "eficiencia"]])

#Graficas

#Horas trabajadas por empleado
plt.bar(df["empleado"], df["horas_trabajadas"], color="skyblue")
plt.title("Horas trabajandas por empleado")
plt.xlabel("Empleado")
plt.ylabel("Horas")
plt.show()

#Eficiencia por empleado
plt.bar(df["empleado"], df["eficiencia"], color="orange")
plt.title("Eficiencia de empleados")
plt.xlabel("Empleado")
plt.ylabel("Eficiencia (tareas/proyectos)")
plt.show()



#Día 241 / 365
"""
Mini reporte de ventas en Excel con gráfico
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

data = {
    "Producto": ["Teclado", "Mouse", "Monitor", "Laptop", "Auriculares"],
    "Ventas": [1200, 800, 1500, 3200, 950]
}

df = pd.DataFrame(data)

ruta = "Mini_Reporte_Ventas.xlsx"
df.to_excel(ruta, index=False, sheet_name="Ventas")

libro = load_workbook(ruta)
hoja = libro["Ventas"]

chart = BarChart()
chart.title = "Ventas por Porducto"
chart.x_axis.title = "Producto"
chart.y_axis.title = "Ventas"

datos = Reference(hoja, min_col=2, min_row=1, max_row=6)
categoria = Reference(hoja, min_col=1, min_row=2, max_row=6)
chart.add_data(datos, titles_from_data=True)
chart.set_categories(categoria)

hoja.add_chart(chart, "D2")

libro.save(ruta)

print("Mini reporte generado con gráfico en 'Mini_Reporte_ventas.xlsx'")


# Día 242 / 365
"""
Análisis de tráfico web y detección de horas pico
"""
import pandas as pd
import numpy as np 
import matplotlib.pyplot as plt
from datetime import datetime, timedelta

# 1.-Creción de Datos de ejemplo
np.random.seed(42)
usuarios = [f"user{i}" for i in range(1, 51)]
fechas = [datetime(2023, 8, 1) + timedelta(minutes=np.random.randint(0, 1440)) for _ in range(200)]
usuarios_random = np.random.choice(usuarios, 200)

df = pd.DataFrame({
    "usuario": usuarios_random,
    "fecha_hora": fechas
})

#Gurdar CSV de ejemplo
df.to_csv("visitas.csv", index=False)

#2.-Leer y procesar CSV
visitas = pd.read_csv("visitas.csv", parse_dates=["fecha_hora"])
visitas["hora"] = visitas["fecha_hora"].dt.hour

#Contar visitas  por hora
visitas_por_hora = visitas.groupby("hora").size()

#Encontrar hora pico
hora_pico = visitas_por_hora.idxmax()
print(f"La hora con más visitas es: {hora_pico}:00 con {visitas_por_hora.max()} visitas.")

#3.-Graficar
plt.bar(visitas_por_hora.index, visitas_por_hora.values, color="skyblue")
plt.title("Visitas por hora")
plt.xlabel("Hora del día")
plt.ylabel("Cantidad de Visitas")
plt.xticks(range(0, 24))
plt.show()


# Día 243 / 365
"""
Sistema simple de recordatorios en consola
"""
import time
from datetime import datetime

#Lista de recordatorios:
recordatorios = [
    ("Reunión con el equipo", "15:30"),
    ("Enviar informe", "16:00"),
    ("Llamar al cliente", "16:15")
]

print("Sistema de recordatorios iniciado...")
print("Recordatorios programados:")
for tarea, hora in recordatorios:
    print(f"- {tarea} a las {hora}")

#Bucle para verificar cada minuto
while True:
    ahora = datetime.now().strftime("%H:%M")
    for tarea, hora in recordatorios:
        if ahora == hora:
            print(f"🔔 ¡Recordatorio! {tarea} (Hora: {hora})")
    time.sleep(60)


# Día 244 / 365
"""
Detector de duplicados en un archivo CSV
"""

import pandas as pd

#Cargar el archivo CSV
#Debes tener una columana con "email" o "id_cliente"
df = pd.read_csv("clientes.csv")

#Mostrar primeras filas
print("Datos originales:")
print(df.head())

#Buscar duplicados en una columna 
duplicados = df[df.duplicated(subset=["email"], keep=False)]

#Mostrar duplicados
print("\nRegistros duplicados encontrados:")
print(duplicados)

#Guardar duplicados en un nuevo archivo
duplicados.to_csv("duplicados.csv", index=False)

print("\nSe ha guardado los duplicados en 'duplicados.csv'")



# Día 245 / 365
"""
Validador de datos de clientes en CSV
"""
import pandas as pd
import re
#Cargar datos desde un archivo CSV
df = pd.read_csv("clientes.csv")

#Funciones de validación
def validar_email(email):
    patron = '^[\w\.-]+@[\w\.-]+\.\w+$'
    return bool(re.match(patron, str(email)))

def validar_telefono(telefono):
    telefono = str(telefono)
    return telefono.isdigit() and 7 <= len(telefono) <= 12

def validar_nombre(nombre):
    return pd.notna(nombre) and str(nombre). strip() != ""

#Crear columnas con validaciones
df["email_valido"] = df["email"].apply(validar_email)
df["telefono_valido"] = df["telefono"].apply(validar_telefono)
df["nombre_valido"] = df["nombre"].apply(validar_nombre)

#Filtrar registros válidos e inválidos
validos = df[df["email_valido"] & df["telefono_valido"] & df["nombre_valido"]]
invalidos = df[~(df["email_valido"] & df["telefono_valido"] & df["nombre_valido"])]

#Guardar en archivos separados
validos.to_csv("clientes_validos.csv", index=False)
invalidos.to_csv("clientes_invalidos.csv", index=False)

print("✅ Clientes válidos guardados en 'clientes_validos.csv'")
print("⚠️ Clientes inválidos guardados en 'clientes_invalidos.csv'")



# ARchivo.csv de ejemplo (clientes.csv)
nombre,email,telefono
Ana Perez,ana.perez@example.com,0987654321
Carlos Ruiz,carlos.ruiz@correo,12345
,juan.gomez@example.com,099112233
Lucía Torres,lucia.torres@example.com,987654321
Pedro López,pedro123@example,09876abc
María Fernández,maria.fernandez@example.com,022334455

